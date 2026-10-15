"""Cross-layer regressions for citation context and precise review comments."""

import pytest
from docx import Document

from app.main import _annotated_rows, _annotate_docx
from app.modules.module2_apa_checker import check_paragraphs
from app.modules.module3_citation_matcher import (
    CitationMatchResult, parse_citations, run_citation_check,
    run_citation_check_paragraphs,
)
from app.modules.prose_extractor import ProseParagraph, extract_prose


def findings(text, is_ref=False):
    para = ProseParagraph(0, "Normal", text, text, None, is_reference_entry=is_ref)
    return check_paragraphs([para], {}, {})


@pytest.mark.parametrize("author", ["Bailey", "Nguyen", "Anderson"])
@pytest.mark.parametrize("prefix", ["", "While ", "Following ", "Although ", "As "])
def test_narrative_citation_is_one_complete_source(author, prefix):
    text = f"{prefix}{author} et al. (2017) describe the findings."
    citations = parse_citations(text)
    assert [(c.authors, c.year) for c in citations] == [([author.lower()], "2017")]
    assert citations[0].raw == f"{author} et al. (2017)"
    assert citations[0].position_in_para == len(prefix)


@pytest.mark.parametrize("author", ["Al Abri", "Van den Berg", "Research and Development Council"])
def test_compound_and_group_names_survive_introductory_words(author):
    result = run_citation_check(
        f"While {author} (2020) found this.", f"{author}. (2020). A study."
    )
    assert result.missing_references == []
    assert result.uncited_references == []


@pytest.mark.parametrize("body", [
    "Prior work (Smith, 2001; Demerouti & Bakker, 2023) supports this.",
    "Prior work (Demerouti & Bakker, 2023; Smith, 2001) supports this.",
    "Prior work (Smith, 2001; Demerouti and Bakker, 2023) supports this.",
    "Prior work (Demerouti & Bakker 2023) supports this.",
    "Prior work (Demerouti & Bakker, 2023, p. 9) supports this.",
    "Prior work (Demerouti & Bakker, 2023 pp. 9–11) supports this.",
    "Prior work (Demerouti & Bakker, 2023, pp. 3, 5; Smith, 2001) supports this.",
    "Prior work (Demerouti & Bakker, 2023, p. iv; Smith, 2001) supports this.",
    "Demerouti and Bakker (2023, p. 9) support this.",
    "Prior work (e.g., Demerouti & Bakker, 2023; Smith, 2001) supports this.",
    "Prior work (Smith, 2001; cf. Demerouti & Bakker, 2023) supports this.",
    "Prior work (Demerouti et al.,, 2023) supports this.",
])
def test_citation_presence_does_not_depend_on_position_or_minor_punctuation(body):
    result = run_citation_check(body, "Demerouti, E., & Bakker, A. B. (2023). A study.")
    assert not any("Demerouti" in u["reference"] for u in result.uncited_references)
    assert any(c.authors[0] == "demerouti" and c.year == "2023" for c in result.citations)


def test_real_missing_uncited_and_year_mismatch_are_retained():
    result = run_citation_check(
        "While Bailey et al. (2017) found this (Demerouti & Bakker, 2023; Smith, 2020).",
        "Bailey, C. (2017). A study.\nDemerouti, E. (2022). A study.\nJones, J. (2020). Another study.",
    )
    assert [u["reference"] for u in result.uncited_references] == ["Jones, J. (2020). Another study."]
    assert [u["citation"] for u in result.year_mismatches] == ["Demerouti & Bakker, 2023"]
    assert [u["citation"] for u in result.missing_references] == ["Smith, 2020"]


def test_grouped_citation_retains_multiple_years_for_one_author():
    result = run_citation_check(
        "Prior work (Smith, 2020, 2021; Jones, 2019) agrees.",
        "Smith, J. (2020). First study.\nSmith, J. (2021). Second study.\nJones, K. (2019). Third study.",
    )
    assert result.uncited_references == []
    assert result.missing_references == []


@pytest.mark.parametrize("text", [
    "Prior work (Battisti et al. 2022; Gerards et al. 2021) agrees.",
    "Prior work (Schaufeli et al., 2002, p. 74) agrees.",
    "Prior work (Soane et al., 2012 p. 9) agrees.",
    "The account refers to p. 9 and pp. 3–5.",
    "The account follows Smith et al. 2021 in its interpretation.",
])
def test_citation_years_and_page_locators_are_not_prose_numbers(text):
    assert not [f for f in findings(text) if f.rule_id in {"MEC002", "MEC003"}]


def test_inline_list_has_list_warning_without_prose_number_warnings():
    fs = findings("The stages are: (1) preparation, (2) action, and (3) reflection.")
    assert any(f.rule_id == "MEC032" for f in fs)
    assert not any(f.rule_id == "MEC003" for f in fs)


def test_actual_sentence_start_and_prose_quantity_still_flagged():
    fs = findings("Prior work (Smith et al. 2021) agrees. 12 interviews revealed 3 themes.")
    assert len([f for f in fs if f.rule_id == "MEC002"]) == 1
    assert len([f for f in fs if f.rule_id == "MEC003"]) == 1


@pytest.mark.parametrize("title", [
    "The impact of working from home: An exploratory study on Hong Kong",
    "Development and application of a new measure: The ISA Engagement Scale",
    "The Employee Engagement Scale: Initial evidence of construct validity and implications for theory and practice",
    "Work in New York City",
    "Research on the United States Department of Education",
])
def test_proper_names_are_not_sufficient_evidence_of_title_case(title):
    assert not [f for f in findings(f"Lee, J. (2020). {title}. Journal, 3, 1–5.", True) if f.rule_id == "REF024"]


@pytest.mark.parametrize("title", [
    "Technology Acceptance In Knowledge-Intensive Firms",
    "Broken Promises: Equity Sensitivity as a Moderator Between Psychological Contract Breach and Employee Attitudes",
])
def test_clear_title_case_is_still_detected(title):
    fs = [f for f in findings(f"Brown, L. (2019). {title}. Organization Studies, 40, 989–1010.", True) if f.rule_id == "REF024"]
    assert len(fs) == 1
    rows = _annotated_rows(fs, CitationMatchResult([], [], [], [], [], [], [], []))
    assert rows[0]["target"] == title


def test_export_comments_point_to_original_numerals_and_title(tmp_path):
    doc = Document()
    text = '  "3 themes" were discussed. We identified 3 themes and later 3 patterns.'
    doc.add_paragraph(text)
    doc.add_paragraph("References", style="Heading 1")
    title = "Technology Acceptance In Knowledge-Intensive Firms"
    doc.add_paragraph(f"Brown, L. (2019). {title}. Organization Studies, 40, 989–1010.")
    ps = extract_prose(doc)
    fs = [f for f in check_paragraphs(ps, {}, {}) if f.rule_id in {"MEC003", "REF024"}]
    rows = _annotated_rows(fs, CitationMatchResult([], [], [], [], [], [], [], []))
    assert len(rows) == 3
    source, output = tmp_path / "source.docx", tmp_path / "reviewed.docx"
    doc.save(source)
    _annotate_docx(str(source), str(output), rows)
    reviewed = Document(output)
    assert [p.text for p in reviewed.paragraphs] == [p.text for p in doc.paragraphs]
    highlighted = [r.text for p in reviewed.paragraphs for r in p.runs if r.font.highlight_color is not None]
    assert highlighted == ["3", "3", title]
    assert len(reviewed.comments) == 3


def test_docx_extraction_to_reference_matching_detects_grouped_sources():
    doc = Document()
    doc.add_paragraph("Introduction", style="Heading 1")
    doc.add_paragraph("This follows earlier findings (Smith, 2001; Demerouti & Bakker, 2023, p. 9).")
    doc.add_paragraph("References", style="Heading 1")
    doc.add_paragraph("Smith, J. (2001). Earlier study.")
    doc.add_paragraph("Demerouti, E., & Bakker, A. B. (2023). Another study.")
    ps = extract_prose(doc)
    result = run_citation_check_paragraphs(ps, "\n".join(p.raw_text for p in ps if p.is_reference_entry))
    assert result.missing_references == []
    assert result.uncited_references == []
