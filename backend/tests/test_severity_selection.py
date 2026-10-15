from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import main
from app.modules.module2_apa_checker import Finding, Severity
from app.modules.module3_citation_matcher import CitationMatchResult


def test_citation_results_have_explicit_consistent_severities():
    result = CitationMatchResult(
        [], [], [{"severity": "error"}, {"severity": "warning"}],
        [{}], [{}], [{}], [{}], [],
    )
    main._decorate_and_sort_citation_issues(result, [])
    assert [i["severity"] for i in result.missing_references] == ["error", "warning"]
    for group in (result.uncited_references, result.year_mismatches,
                  result.spelling_mismatches, result.co_author_only_matches):
        assert [i["severity"] for i in group] == ["warning"]


@pytest.fixture
def review_client(monkeypatch):
    async def access():
        pass

    main.app.dependency_overrides[main.require_review_access] = access
    findings = [
        Finding("TEST", severity, idx, f"{severity.value} finding", excerpt=f"Target {idx}")
        for idx, severity in enumerate(Severity) if severity.value != "info"
    ]
    findings.append(Finding("TEST", Severity.INFO, 3, "info finding", excerpt="Target 3"))
    monkeypatch.setattr(main, "check_paragraphs", lambda *args: findings)
    try:
        yield TestClient(main.app)
    finally:
        main.app.dependency_overrides.clear()


@pytest.mark.parametrize("selected,expected", [
    (None, {"error", "warning", "suggestion", "info"}),
    (["error", "warning"], {"error", "warning"}),
    (["error"], {"error"}),
    (["warning"], {"warning"}),
    (["suggestion"], {"suggestion"}),
    (["info"], {"info"}),
])
def test_annotated_download_contains_only_selected_severities(review_client, selected, expected):
    doc = Document()
    for idx in range(4):
        doc.add_paragraph(f"Target {idx}")
    stream = BytesIO()
    doc.save(stream)
    data = {"severities": selected} if selected is not None else {}
    response = review_client.post(
        "/api/check/docx/annotated", data=data,
        files={"file": ("article.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    reviewed = Document(BytesIO(response.content))
    assert [p.text for p in reviewed.paragraphs] == [p.text for p in doc.paragraphs]
    assert {c.text.split("Severity: ")[1].split("\n")[0].lower() for c in reviewed.comments} == expected
    assert len(reviewed.comments) == len(expected)


def test_annotated_download_rejects_unknown_severity(review_client):
    response = review_client.post(
        "/api/check/docx/annotated", data={"severities": ["urgent"]},
        files={"file": ("article.docx", b"not needed for validation")},
    )
    assert response.status_code == 422


@pytest.mark.parametrize("severity", ["error", "warning"])
def test_download_filters_citation_comments_with_apa_comments(review_client, monkeypatch, severity):
    result = CitationMatchResult(
        [], [], [{"citation": "(Missing, 2020)", "paragraph_index": 0,
                  "severity": "error", "message": "Missing source"}], [],
        [{"citation": "(Year, 2020)", "paragraph_index": 1, "message": "Check year"}],
        [], [], [],
    )
    monkeypatch.setattr(main, "run_citation_check_paragraphs", lambda *args, **kwargs: result)
    doc = Document()
    for idx, citation in enumerate(["(Missing, 2020)", "(Year, 2020)", "", ""]):
        doc.add_paragraph(f"Target {idx} {citation}")
    stream = BytesIO()
    doc.save(stream)
    response = review_client.post(
        "/api/check/docx/annotated",
        data={"severities": [severity], "reference_text": "References supplied"},
        files={"file": ("article.docx", stream.getvalue())},
    )
    assert response.status_code == 200
    comments = list(Document(BytesIO(response.content)).comments)
    assert len(comments) == 2
    assert all(f"Severity: {severity.upper()}" in c.text for c in comments)
    assert any(c.text.startswith("APA:") for c in comments)
    assert any(c.text.startswith("Citation:") for c in comments)
